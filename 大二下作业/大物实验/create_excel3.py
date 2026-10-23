import os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill

# 设置文件路径
file_path = '悬挂法测量数据_新版.xlsx'

# 加载现有工作簿
try:
    wb = load_workbook(file_path)
except FileNotFoundError:
    print(f"错误：找不到文件 {file_path}，请先运行create_excel.py创建文件")
    exit(1)

# 创建表3
ws3 = wb.create_sheet('表3')

# 设置标题行
ws3.merge_cells('A1:D1')
ws3['A1'] = '表3 被测试样的参数与共振基频记录'
ws3['A1'].alignment = Alignment(horizontal='center')
ws3['A1'].font = Font(bold=True)

# 设置室温和日期行
ws3.merge_cells('A2:D2')
ws3['A2'] = '室温：________      实验日期：________'
ws3['A2'].alignment = Alignment(horizontal='center')

# 创建表头和数据
data = [
    ['测试样品材质', '黄铜圆柱体样', '不锈钢圆柱体样'],
    ['样品长度 L (×10⁻²m)', 16.0, 16.0],
    ['样品直径 d (×10⁻³m)', 6.00, 6.00],
    ['样品质量 m (×10⁻³kg)', 36.60, 35.00],
    ['悬挂共振频率 f(Hz)', '', ''],
    ['支撑共振频率 f(Hz)', '', '不测']
]

# 填充数据
for i, row_data in enumerate(data):
    for j, value in enumerate(row_data):
        ws3.cell(row=i+3, column=j+1, value=value)
        
# 添加边框
thin_border = Side(border_style='thin')
border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

# 设置背景色为浅绿色
light_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

# 调整列宽
column_letters = ['A', 'B', 'C', 'D']
for col_letter in column_letters:
    ws3.column_dimensions[col_letter].width = 20

# 为所有单元格添加边框和对齐方式
for row in range(1, len(data) + 3):  # +3 因为有标题行和室温日期行
    for col in range(1, 4):  # 只有3列有数据
        cell = ws3.cell(row=row, column=col)
        if not isinstance(cell, type(None)):
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            # 除了标题行和室温日期行外的所有行添加背景色
            if row > 2:
                cell.fill = light_green

# 保存修改后的文件
wb.save(file_path)

print(f'已在Excel文件添加表3：{os.path.abspath(file_path)}') 