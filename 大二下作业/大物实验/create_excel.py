import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill

# 创建数据 - 按照正确的行列结构
data = [
    ['测量组别', '1', '2', '3', '4', '5', '6'],
    ['左悬挂点位置', 5.84, 15.84, 25.84, 35.84, 45.84, 55.84],
    ['右悬挂点位置', 154.16, 144.16, 134.16, 124.16, 114.16, 104.16],
    ['黄铜棒共振频率', '', '', '', '/', '', ''],
    ['不锈钢棒共振频率', '', '', '', '/', '', '']
]

# 设置文件路径
file_path = '悬挂法测量数据_新版.xlsx'

# 创建新的工作簿
wb = Workbook()
ws = wb.active
ws.title = '表1'

# 将数据写入工作表
for i, row_data in enumerate(data):
    for j, value in enumerate(row_data):
        ws.cell(row=i+1, column=j+1, value=value)

# 插入标题行
ws.insert_rows(1)
ws.merge_cells('A1:G1')
ws['A1'] = '表1 悬挂法测黄铜棒和不锈钢棒的共振频率数据记录'
ws['A1'].alignment = Alignment(horizontal='center')
ws['A1'].font = Font(bold=True)

# 添加边框
thin_border = Side(border_style='thin')
border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

# 设置背景色为浅绿色
light_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

# 调整列宽和设置样式
column_letters = ['A', 'B', 'C', 'D', 'E', 'F', 'G']
for col_letter in column_letters:
    ws.column_dimensions[col_letter].width = 15

# 为所有单元格添加边框和对齐方式
for row in range(1, len(data) + 2):  # +2 因为添加了标题行
    for col in range(1, len(data[0]) + 1):
        cell = ws.cell(row=row, column=col)
        if not isinstance(cell, type(None)):  # 检查单元格是否为None
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
            # 除了标题行外的所有行添加背景色
            if row > 1:
                cell.fill = light_green

# 保存修改后的文件
wb.save(file_path)

print(f'Excel文件已创建：{os.path.abspath(file_path)}') 